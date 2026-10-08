"""База контактов для точечной рассылки: content/outreach/data/*.json → contacts.xlsx + contacts.csv.

Запуск из корня репозитория: python3 content/outreach/build-contacts.py   (нужен openpyxl)
Каждый JSON — список объектов: name, category, format, city, dates, description, url, emails[], phones[], socials[],
contact_person, audience, priority (A/B/C), source. Сегмент берётся из имени файла (SEGMENTS).
"""
import csv
import hashlib
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DIR = Path(__file__).resolve().parent
SEGMENTS = {  # файл → сегмент (порядок = порядок в таблице)
    "conferences": "Конференции и олимпиады",
    "forums": "Форумы, конгрессы, выставки",
    "media": "СМИ и журналы",
    "communities": "Онлайн-сообщества",
    "orgs": "Студенческие организации и вузы",
    "partners": "Партнёры с общей аудиторией",
}
MERGED = []  # (оставили, слили) — для проверки
EVENT_SEGMENTS = {"Конференции и олимпиады", "Форумы, конгрессы, выставки"}
STATUSES = ["Не начато", "Написали", "Позвонили", "Ответили", "Договорились", "Отказ", "Нет ответа"]
COLUMNS = [  # заголовок, ключ, ширина
    ("№", "n", 5), ("Название", "name", 34), ("Сегмент", "segment", 20), ("Категория", "category", 18),
    ("Формат", "format", 10), ("Город", "city", 14), ("Даты", "dates", 18), ("Кто это (одно предложение)", "description", 60),
    ("Ссылка", "url", 34), ("Email", "emails", 32), ("Телефон", "phones", 22), ("Соцсети / мессенджеры", "socials", 40),
    ("Контактное лицо", "contact_person", 26), ("Аудитория", "audience", 18), ("Приоритет", "priority", 10),
    ("Где нашли контакты", "source", 30), ("Статус", "status", 14), ("Кто ведёт", "owner", 12), ("Комментарий", "comment", 30),
]


def norm_url(u):
    u = (u or "").strip().lower()
    u = re.sub(r"^https?://(www\.)?", "", u).rstrip("/")
    return u


def words(name):
    return {w for w in re.findall(r"\w+", name.lower()) if len(w) > 2 and w not in {"vk", "telegram", "канал"}}


def same(a, b):
    """Дубль: одно мероприятие с тем же названием в «кавычках», один и тот же проект под тем же именем
    (Telegram и VK), или одна ссылка при похожих названиях — общая страница вуза сама по себе не повод склеивать."""
    if (a["segment"] in EVENT_SEGMENTS) != (b["segment"] in EVENT_SEGMENTS) and "Студенческие организации и вузы" in (a["segment"], b["segment"]):
        return False  # СНО и его конференция — разные поводы написать
    wa, wb = words(a["name"]), words(b["name"])
    sim = len(wa & wb) / max(1, min(len(wa), len(wb)))
    qa, qb = re.findall(r"«([^»]{4,})»", a["name"]), re.findall(r"«([^»]{4,})»", b["name"])
    if a["segment"] in EVENT_SEGMENTS and b["segment"] in EVENT_SEGMENTS and qa and qb:
        return re.sub(r"\W", "", qa[0].lower()) == re.sub(r"\W", "", qb[0].lower()) and sim >= 0.4
    base = lambda r: re.sub(r"\(.*?\)|[^\w]", "", r["name"].lower())
    if a["segment"] == b["segment"] and base(a) and base(a) == base(b):
        return True
    ua, ub = norm_url(a.get("url")), norm_url(b.get("url"))
    return bool(ua) and ua == ub and sim >= 0.75


def clean_list(v):
    if not v:
        return []
    if isinstance(v, str):
        v = re.split(r"[;,\n]\s*", v)
    out = []
    for x in v:
        x = str(x).strip()
        if x and x.lower() not in {o.lower() for o in out}:
            out.append(x)
    return out


def load():
    rows = []
    for key, seg in SEGMENTS.items():
        f = DIR / "data" / f"{key}.json"
        if not f.exists():
            continue
        for r in json.loads(f.read_text("utf-8")):
            name = (r.get("name") or "").strip()
            if not name:
                continue
            r = {k: (v.strip() if isinstance(v, str) else v) for k, v in r.items()}
            r["name"] = name
            r["segment"] = seg
            for k in ("emails", "phones", "socials"):
                r[k] = clean_list(r.get(k))
            r["priority"] = (r.get("priority") or "C").strip()[:1].upper()
            first = next((o for o in rows if same(o, r)), None)
            if first:  # дубль (часто из соседнего сегмента) — сливаем контакты в первую запись
                extra = [r["url"]] if norm_url(r.get("url")) not in ("", norm_url(first.get("url"))) else []
                for k in ("emails", "phones", "socials"):
                    first[k] = clean_list(first[k] + r[k] + (extra if k == "socials" else []))
                MERGED.append((first["name"], r["name"]))
                continue
            rows.append(r)
    order = {s: i for i, s in enumerate(SEGMENTS.values())}
    rows.sort(key=lambda r: (order[r["segment"]], r["priority"], r["name"].lower()))
    for i, r in enumerate(rows, 1):
        r["n"] = i
        r["id"] = hashlib.sha1(r["name"].encode()).hexdigest()[:8]  # стабильный ключ для рассылки (mailing/letters.json)
        r.setdefault("status", "Не начато")
    return rows


def cell_value(r, key):
    v = r.get(key, "")
    if isinstance(v, list):
        return "\n".join(v)
    return v if v is not None else ""


def build(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Контакты"
    head_fill = PatternFill("solid", fgColor="0D5C55")
    prio_fill = {"A": PatternFill("solid", fgColor="DCEAE5"), "B": PatternFill("solid", fgColor="FFF7DA")}
    ws.append([c[0] for c in COLUMNS])
    for i, (_, _, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    keys = [c[1] for c in COLUMNS]
    for r in rows:
        ws.append([cell_value(r, k) for k in keys])
        row = ws.max_row
        for i, k in enumerate(keys, 1):
            c = ws.cell(row=row, column=i)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if k == "url" and r.get("url", "").startswith("http"):
                c.hyperlink = r["url"]
                c.font = Font(color="0D5C55", underline="single")
            if k == "priority" and r["priority"] in prio_fill:
                c.fill = prio_fill[r["priority"]]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    col = get_column_letter(keys.index("status") + 1)
    dv.add(f"{col}2:{col}{max(ws.max_row, 2) + 200}")

    # Сводка: сколько записей и с какими контактами в каждом сегменте
    s = wb.create_sheet("Сводка")
    s.append(["Сегмент", "Всего", "Приоритет A", "С email", "С телефоном", "С соцсетями", "Без контактов"])
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    for seg in list(SEGMENTS.values()) + ["Итого"]:
        rs = rows if seg == "Итого" else [r for r in rows if r["segment"] == seg]
        s.append([seg, len(rs), sum(r["priority"] == "A" for r in rs), sum(bool(r["emails"]) for r in rs),
                  sum(bool(r["phones"]) for r in rs), sum(bool(r["socials"]) for r in rs),
                  sum(not (r["emails"] or r["phones"] or r["socials"]) for r in rs)])
    s[s.max_row][0].font = Font(bold=True)
    for i, w in enumerate([34, 8, 13, 10, 13, 13, 15], 1):
        s.column_dimensions[get_column_letter(i)].width = w
    s.append([])
    s.append(["Приоритет A — большая профильная аудитория и есть прямой контакт; B — стоит написать; C — по остаточному принципу."])
    s.append(["Контакты собраны из открытых официальных источников (колонка «Где нашли контакты»). Перед отправкой проверьте, что адрес ещё актуален."])
    s.append(["Тексты писем и сообщений — templates.md, презентация — help-me-doctor-presentation.pdf."])
    wb.save(DIR / "contacts.xlsx")

    with open(DIR / "contacts.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([c[0] for c in COLUMNS])
        for r in rows:
            w.writerow([cell_value(r, k).replace("\n", "; ") if isinstance(cell_value(r, k), str) else cell_value(r, k) for k in keys])


def export_recipients(rows):
    """Адресаты рассылки — все, у кого есть email. Тексты писем к ним лежат в mailing/letters.json."""
    keep = ("id", "segment", "category", "name", "description", "url", "emails", "contact_person", "city", "dates", "priority")
    out = [{k: r.get(k, "") for k in keep} for r in rows if r["emails"]]
    (DIR / "mailing").mkdir(exist_ok=True)
    (DIR / "mailing" / "recipients.json").write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", "utf-8")
    return out


if __name__ == "__main__":
    rows = load()
    build(rows)
    print(f"✓ mailing/recipients.json: {len(export_recipients(rows))} адресатов с email")
    by = {}
    for r in rows:
        by[r["segment"]] = by.get(r["segment"], 0) + 1
    if "-v" in __import__("sys").argv:
        for a, b in MERGED:
            print(f"  = {a}\n    ← {b}")
    print(f"✓ contacts.xlsx: {len(rows)} записей", by)
